"""Run a folder of documents through read -> extract -> check, and export the results.

Text is cached in <out>/../_text so OCR runs once per document, whichever extractor is used.
Every document lands in invoices.xlsx: 'Ready' when every check passes, 'Needs review' otherwise.
"""
from __future__ import annotations

import hashlib
import json
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from . import extract_rules, ingest, validate
from .ingest import read_text
from .schema import Invoice

DOC_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
_READER = Path(ingest.__file__).read_bytes()


def _text(doc: Path, cache_dir: Path) -> tuple[str, str]:
    """Read a document's text once and reuse it. The cache key is a hash of the file's CONTENT plus the reading
    code, never the file's name: two test sets both have a doc_002, and a name-keyed cache once served one set's
    text for the other. Because the reading code is part of the key, a fix to it is never hidden by old text."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha1(_READER + doc.read_bytes()).hexdigest()[:16]
    tf, mf = cache_dir / f"{key}.txt", cache_dir / f"{key}.method"
    if tf.exists() and mf.exists():
        return tf.read_text(encoding="utf-8"), mf.read_text(encoding="utf-8")
    text, method = read_text(doc)
    tf.write_text(text, encoding="utf-8")
    mf.write_text(method, encoding="utf-8")
    return text, method


def run(docs_dir: str, out_dir: str, backend: str = "rules", model: str | None = None,
        resume: bool = False, limit: int | None = None) -> list[dict]:
    """resume=True keeps the answers from an earlier, interrupted run (rerunning the checks on them) and reads only
    the missing or failed documents. limit=N reads only the first N documents (handy for a quick look at the AI)."""
    docs = sorted(p for p in Path(docs_dir).iterdir() if p.suffix.lower() in DOC_SUFFIXES)[:limit]
    out = Path(out_dir)
    (out / "pred").mkdir(parents=True, exist_ok=True)
    if backend == "rules":
        extractor = extract_rules.extract
    else:
        from . import extract_llm  # only needs Ollama when this backend is used
        extractor = (lambda t: extract_llm.extract(t, model)) if model else extract_llm.extract
    rows = []
    for doc in docs:
        done = out / "pred" / f"{doc.stem}.json"
        text, method = _text(doc, out.parent / "_text")
        if resume and done.exists():
            rec = json.loads(done.read_text(encoding="utf-8"))
            if not rec["error"]:  # a document that failed (say, Ollama was stopped mid-run) is read again
                if rec["invoice"]:  # keep the answer, but rerun today's checks on it, so a new check needs no new AI run
                    rec["issues"] = validate.check(Invoice.model_validate(rec["invoice"]), text)
                    done.write_text(json.dumps(rec, indent=1, ensure_ascii=False), encoding="utf-8")
                rows.append(rec)
                continue
        t0 = time.perf_counter()
        try:
            inv, error = extractor(text), None
            issues = validate.check(inv, text)
        except Exception as e:  # a failed read is a review item, never a silent skip
            inv, error, issues = None, f"{type(e).__name__}: {e}", ["extraction failed"]
        rec = {"doc_id": doc.stem, "file": doc.name, "method": method, "backend": backend, "model": model,
               "seconds": round(time.perf_counter() - t0, 2), "error": error, "issues": issues,
               "invoice": inv.model_dump() if inv else None}
        (out / "pred" / f"{doc.stem}.json").write_text(json.dumps(rec, indent=1, ensure_ascii=False), encoding="utf-8")
        rows.append(rec)
        print(f"{doc.name:12} {method:8} {'REVIEW' if issues else 'ready ':6} {rec['seconds']:6.2f}s  "
              f"{'; '.join(issues)[:90]}")
    export_xlsx(rows, out / "invoices.xlsx")
    return rows


HEAD = ["doc_id", "vendor_name", "invoice_number", "issue_date", "due_date", "currency", "subtotal", "discount",
        "shipping", "tax", "total", "items", "read_by", "issues"]


def export_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ready, review = wb.active, wb.create_sheet("Needs review")
    ready.title = "Ready"
    lines = wb.create_sheet("Line items")
    for ws in (ready, review):
        ws.append(HEAD)
    lines.append(["doc_id", "description", "quantity", "unit_price", "amount"])
    for r in rows:
        inv = r["invoice"] or {}
        vals = [r["doc_id"]] + [inv.get(k) for k in HEAD[1:11]] + [len(inv.get("line_items") or []), r["method"],
                                                                     "; ".join(r["issues"])]
        (review if r["issues"] else ready).append(vals)
        for it in inv.get("line_items") or []:
            lines.append([r["doc_id"], it["description"], it["quantity"], it["unit_price"], it["amount"]])
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        # Start the cursor below the frozen header. With it left on A1 (inside the frozen row), Excel scrolls the
        # lower pane up to show row 1 again, so the header appears twice.
        ws.sheet_view.selection[0].activeCell = "A2"
        ws.sheet_view.selection[0].sqref = "A2"
        # Size each column to its longest value, so nothing is cut off ("Harbor an…") when the file opens.
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max(longest + 2, 9), 60)
    wb.save(path)


def load_preds(pred_dir: str) -> dict[str, dict]:
    return {p.stem: json.loads(p.read_text(encoding="utf-8")) for p in Path(pred_dir, "pred").glob("*.json")}


def load_truth(truth_dir: str) -> dict[str, Invoice]:
    return {p.stem: Invoice.model_validate_json(p.read_text(encoding="utf-8")) for p in Path(truth_dir).glob("*.json")}
